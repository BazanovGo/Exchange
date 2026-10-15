"""Trade ledger of the selected configurations: every trade in order with its result.

For each instrument the stage-3 finalist (best IS plateau, ``tables/finalists.csv``)
is re-simulated on the 5m candles. Each trade gets its entry and exit contract
and raw price (the bar open the order fills at), the number of rolls it was
carried through, its PnL in USD and ticks for one contract, and the running
total. PnL comes from the held-contract path: within one contract it equals
direction x (exit - entry) x contract size; a trade carried through a roll
switches contracts at the overlap price, so its raw prices belong to
different contracts.
"""

from __future__ import annotations

from pathlib import Path

import numpy as np
import pandas as pd

from src.ma_slope import finalist as F
from src.ma_slope.data import TICK_SIZES, TICKERS, load_instrument
from src.ma_slope.grid import EXIT_MODELS, Slope
from src.ma_slope.presentation import NAMES, TABLES, grid_meta

LEDGER_CSV = TABLES / "trade_ledger.csv"
LEDGER_XLSX = TABLES / "trade_ledger.xlsx"
ORDER = ("ng", "br", "gold", "silv", "plt", "pld", "copper", "wti")
assert set(ORDER) == set(TICKERS)


def finalist_config(ticker: str, stage: str = "stage3") -> F.Config:
    fz = pd.read_csv(TABLES / "finalists.csv")
    row = fz[(fz["stage"] == stage) & (fz["tag"] == "overall") & (fz["ticker"] == ticker)].iloc[0]
    slope = Slope(**grid_meta(ticker, stage)["slopes"][f"{row['method']}/{row['model']}"])
    return F.Config(row["fast_type"], int(row["fast_period"]), row["slow_type"], int(row["slow_period"]), slope,
                    EXIT_MODELS.index(row["model"]))


def ledger(ticker: str, stage: str = "stage3") -> tuple[pd.DataFrame, F.Config, pd.Timestamp]:
    """All closed trades of the instrument's finalist, in order, with prices, contracts and PnL."""
    cfg = finalist_config(ticker, stage)
    ins = load_instrument(ticker)
    tr = F.trades(ins, cfg)
    e = ins.index.get_indexer(tr["entry_time"])
    x = ins.index.get_indexer(tr["exit_time"])
    raw_open = ins.open / ins.factor
    direction = tr["direction"].to_numpy()
    changes = np.r_[0, np.cumsum(ins.secid[1:] != ins.secid[:-1])]
    pnl = tr["pnl"].to_numpy()

    # The ledger must reproduce the simulator: PnL is the held-contract path move, and inside one
    # contract that is the raw price move times the contract size.
    np.testing.assert_allclose(pnl, direction * (ins.path[x] - ins.path[e]) * ins.multiplier, atol=1e-6)
    same = changes[x] == changes[e]
    np.testing.assert_allclose(pnl[same], (direction * (raw_open[x] - raw_open[e]) * ins.multiplier)[same],
                               atol=1e-6)

    tick_value = TICK_SIZES[ticker] * ins.multiplier
    out = pd.DataFrame({
        "instrument": NAMES[ticker],
        "ticker": ticker,
        "trade": np.arange(1, len(tr) + 1),
        "sample": tr["sample"].str.upper().to_numpy(),
        "direction": np.where(direction == 1, "Long", "Short"),
        "entry_time": tr["entry_time"].to_numpy(),
        "entry_contract": ins.secid[e],
        "entry_price": np.round(raw_open[e], 6),
        "exit_time": tr["exit_time"].to_numpy(),
        "exit_contract": ins.secid[x],
        "exit_price": np.round(raw_open[x], 6),
        "rolls": changes[x] - changes[e],
        "hours": tr["hours"].to_numpy(),
        "pnl_usd": np.round(pnl, 4),
        "pnl_ticks": np.round(pnl / tick_value, 2),
        "outcome": np.where(pnl > 0, "profit", np.where(pnl < 0, "loss", "zero")),
        "cum_pnl_usd": np.round(np.cumsum(pnl), 4),
    })
    out.attrs["multiplier"] = ins.multiplier
    return out, cfg, ins.split_time


def all_ledgers(stage: str = "stage3") -> dict[str, tuple[pd.DataFrame, F.Config, pd.Timestamp]]:
    return {t: ledger(t, stage) for t in ORDER}


def write_csv(ledgers, path: Path = LEDGER_CSV) -> Path:
    frames = [df.assign(config=cfg.label) for df, cfg, _ in ledgers.values()]
    pd.concat(frames, ignore_index=True).to_csv(path, index=False)
    return path


SHEET = {"ng": "NG", "br": "BR", "gold": "GOLD", "silv": "SILV", "plt": "PLT", "pld": "PLD", "copper": "COPPER",
         "wti": "WTI"}
PRICE_FORMAT = {"ng": "0.000", "br": "0.00", "gold": "0.0", "silv": "0.00", "plt": "0.0", "pld": "0.00",
                "copper": "#,##0", "wti": "0.00"}
USD = "#,##0.00;[Red]-#,##0.00;0.00"
TICKS = "#,##0.0;[Red]-#,##0.0;0.0"
TIME = "yyyy-mm-dd hh:mm"
FIRST = 9  # first data row on an instrument sheet
NOTE = ("PnL — на 1 контракт, без комиссий и проскальзывания. Вход и выход — по Open свечи после сигнала; цены — "
        "сырые цены контракта. Сделка, прошедшая через ролл, переходит в новый контракт по цене перекрытия, поэтому "
        "её цены входа и выхода относятся к разным контрактам. Выборка (IS/OOS) — по времени выхода.")


def write_excel(ledgers, path: Path = LEDGER_XLSX) -> Path:
    """Workbook: summary (formulas), all trades by exit time, one sheet per instrument in trade order."""
    from openpyxl import Workbook
    from openpyxl.comments import Comment
    from openpyxl.formatting.rule import FormulaRule
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    base, bold, title_font = Font(name="Arial", size=10), Font(name="Arial", size=10, bold=True), Font(
        name="Arial", size=12, bold=True)
    blue, link = Font(name="Arial", size=10, color="0000FF"), Font(name="Arial", size=10, color="008000")
    muted = Font(name="Arial", size=9, italic=True, color="595959")
    head_fill = PatternFill("solid", fgColor="E7E6E1")
    head_border = Border(bottom=Side(style="thin", color="7F7F7F"))
    wrap = Alignment(wrap_text=True, vertical="top")
    green_font, red_font = Font(name="Arial", size=10, color="1B7A3D"), Font(name="Arial", size=10, color="C00000")

    def header(ws, row, names, widths):
        for j, (name, width) in enumerate(zip(names, widths, strict=True), start=1):
            c = ws.cell(row, j, name)
            c.font, c.fill, c.border, c.alignment = bold, head_fill, head_border, Alignment(wrap_text=True,
                                                                                          vertical="center")
            ws.column_dimensions[get_column_letter(j)].width = width
        ws.row_dimensions[row].height = 30

    def put(ws, row, col, value, fmt=None, font=base):
        c = ws.cell(row, col, value)
        c.font = font
        if fmt:
            c.number_format = fmt
        return c

    def outcome_colours(ws, col, first, last):
        rng = f"{col}{first}:{col}{last}"
        ws.conditional_formatting.add(rng, FormulaRule(formula=[f'{col}{first}="Прибыль"'], font=green_font))
        ws.conditional_formatting.add(rng, FormulaRule(formula=[f'{col}{first}="Убыток"'], font=red_font))

    wb = Workbook()
    summary = wb.active
    summary.title = "Сводка"
    everything = wb.create_sheet("Все сделки")
    rows_of = {}

    # One sheet per instrument: every trade in order, running total, peak, drawdown and loss streak as formulas.
    for t, (df, cfg, split) in ledgers.items():
        ws = wb.create_sheet(SHEET[t])
        put(ws, 1, 1, f"{NAMES[t]}: {cfg.label} — все сделки по порядку", font=title_font)
        put(ws, 2, 1, "Размер контракта, USD за 1,0 цены")
        put(ws, 2, 2, df.attrs.get("multiplier", None), font=blue)
        put(ws, 3, 1, "Шаг цены (тик)")
        put(ws, 3, 2, TICK_SIZES[t], font=blue)
        put(ws, 4, 1, "Стоимость тика, USD")
        put(ws, 4, 2, "=B2*B3", "0.00")
        put(ws, 5, 1, "Граница IS/OOS")
        put(ws, 5, 2, split.to_pydatetime(), TIME)
        ws.cell(2, 3, "спецификация MOEX (src/ma_slope/data.py)").font = muted
        ws.cell(5, 3, "первые 70 % 5-минутных баров — IS").font = muted
        ws.merge_cells("A6:R6")
        put(ws, 6, 1, NOTE, font=muted).alignment = wrap
        ws.row_dimensions[6].height = 40
        names = ["№", "Выборка", "Направление", "Вход: время", "Вход: контракт", "Вход: цена", "Выход: время",
                 "Выход: контракт", "Выход: цена", "Роллов внутри сделки", "Длительность, ч", "PnL, $", "Исход",
                 "PnL, тиков", "Накопленный PnL, $", "Пик накопленного, $", "Просадка от пика, $",
                 "Убытков подряд"]
        header(ws, FIRST - 1, names, [7, 8, 11, 17, 12, 11, 17, 12, 11, 9, 10, 11, 10, 10, 13, 13, 12, 9])
        ws.cell(FIRST - 1, 12).comment = Comment("Результат сделки из симулятора исследования (держимый контракт, "
                                                 "1 контракт). Совпадает с tables/finalist_trades.", "ledger")
        pf = PRICE_FORMAT[t]
        last = FIRST + len(df) - 1
        for i, r in enumerate(df.itertuples(index=False)):
            row = FIRST + i
            put(ws, row, 1, int(r.trade))
            put(ws, row, 2, r.sample)
            put(ws, row, 3, r.direction)
            put(ws, row, 4, pd.Timestamp(r.entry_time).to_pydatetime(), TIME)
            put(ws, row, 5, r.entry_contract)
            put(ws, row, 6, float(r.entry_price), pf)
            put(ws, row, 7, pd.Timestamp(r.exit_time).to_pydatetime(), TIME)
            put(ws, row, 8, r.exit_contract)
            put(ws, row, 9, float(r.exit_price), pf)
            put(ws, row, 10, int(r.rolls))
            put(ws, row, 11, float(r.hours), "0.0")
            put(ws, row, 12, float(r.pnl_usd), USD)
            put(ws, row, 13, f'=IF(L{row}>0,"Прибыль",IF(L{row}<0,"Убыток","Ноль"))')
            put(ws, row, 14, f"=L{row}/$B$4", TICKS)
            put(ws, row, 15, f"=L{row}" if row == FIRST else f"=O{row - 1}+L{row}", USD)
            put(ws, row, 16, f"=MAX(0,O{row})" if row == FIRST else f"=MAX(P{row - 1},O{row})", USD)
            put(ws, row, 17, f"=P{row}-O{row}", USD)
            put(ws, row, 18, f"=IF(L{row}<0,1,0)" if row == FIRST else f"=IF(L{row}<0,R{row - 1}+1,0)")
        outcome_colours(ws, "M", FIRST, last)
        ws.freeze_panes = f"B{FIRST}"
        ws.auto_filter.ref = f"A{FIRST - 1}:R{last}"
        rows_of[t] = (FIRST, last)

    # All trades of all instruments in the order they closed; PnL links to the instrument sheets.
    allt = pd.concat([df.assign(sheet_row=np.arange(FIRST, FIRST + len(df))) for df, _, _ in ledgers.values()])
    allt = allt.sort_values(["exit_time", "entry_time", "ticker"], kind="stable").reset_index(drop=True)
    put(everything, 1, 1, "Все сделки всех инструментов в порядке закрытия (по 1 контракту каждого)", font=title_font)
    everything.merge_cells("A2:K2")
    put(everything, 2, 1, "Накопленный итог — сумма по всем инструментам в долларах США; WTI торгуется на MOEX "
        "с июля 2026 года и приведён справочно. PnL берётся ссылкой с листа инструмента.", font=muted).alignment = wrap
    everything.row_dimensions[2].height = 28
    header(everything, 4, ["№", "Инструмент", "№ сделки в инструменте", "Выборка", "Направление", "Вход: время",
                           "Выход: время", "Длительность, ч", "PnL, $", "Исход", "Накопленный PnL, $"],
           [8, 13, 11, 8, 11, 17, 17, 10, 11, 10, 14])
    for i, r in enumerate(allt.itertuples(index=False)):
        row = 5 + i
        put(everything, row, 1, i + 1)
        put(everything, row, 2, r.instrument)
        put(everything, row, 3, int(r.trade))
        put(everything, row, 4, r.sample)
        put(everything, row, 5, r.direction)
        put(everything, row, 6, pd.Timestamp(r.entry_time).to_pydatetime(), TIME)
        put(everything, row, 7, pd.Timestamp(r.exit_time).to_pydatetime(), TIME)
        put(everything, row, 8, float(r.hours), "0.0")
        put(everything, row, 9, f"='{SHEET[r.ticker]}'!L{r.sheet_row}", USD, font=link)
        put(everything, row, 10, f'=IF(I{row}>0,"Прибыль",IF(I{row}<0,"Убыток","Ноль"))')
        put(everything, row, 11, f"=I{row}" if row == 5 else f"=K{row - 1}+I{row}", USD)
    last_all = 4 + len(allt)
    outcome_colours(everything, "J", 5, last_all)
    everything.freeze_panes = "B5"
    everything.auto_filter.ref = f"A4:K{last_all}"

    # Summary: every figure is a formula over the instrument sheets.
    put(summary, 1, 1, "Все сделки лучших конфигураций (этап 3): сводка по инструментам", font=title_font)
    summary.merge_cells("A2:Q2")
    put(summary, 2, 1, "Конфигурация каждого инструмента — лучшая по IS-плато (tables/finalists.csv). " + NOTE,
        font=muted).alignment = wrap
    summary.row_dimensions[2].height = 52
    cols = ["Инструмент", "Конфигурация", "Сделок", "Прибыльных", "Убыточных", "Доля прибыльных", "Итого PnL, $",
            "PnL IS, $", "PnL OOS, $", "Средняя сделка, $", "Средняя сделка, тиков", "Средняя прибыль, $",
            "Средний убыток, $", "Лучшая сделка, $", "Худшая сделка, $", "Макс. убытков подряд",
            "Макс. просадка, $"]
    header(summary, 4, cols, [13, 44, 9, 11, 11, 11, 12, 12, 12, 12, 12, 12, 12, 12, 12, 11, 12])
    row = 5
    for t, (_, cfg, _) in ledgers.items():
        s = f"'{SHEET[t]}'!"
        a, b = rows_of[t]
        pnl, smp = f"{s}$L${a}:$L${b}", f"{s}$B${a}:$B${b}"
        put(summary, row, 1, NAMES[t], font=bold)
        put(summary, row, 2, cfg.label)
        put(summary, row, 3, f"=COUNT({pnl})", "#,##0", link)
        put(summary, row, 4, f'=COUNTIF({pnl},">0")', "#,##0", link)
        put(summary, row, 5, f'=COUNTIF({pnl},"<0")', "#,##0", link)
        put(summary, row, 6, f"=IF(C{row}>0,D{row}/C{row},0)", "0.0%")
        put(summary, row, 7, f"=SUM({pnl})", USD, link)
        put(summary, row, 8, f'=SUMIF({smp},"IS",{pnl})', USD, link)
        put(summary, row, 9, f'=SUMIF({smp},"OOS",{pnl})', USD, link)
        put(summary, row, 10, f"=IF(C{row}>0,G{row}/C{row},0)", USD)
        put(summary, row, 11, f"=IF(C{row}>0,J{row}/{s}$B$4,0)", TICKS, link)
        put(summary, row, 12, f'=IFERROR(AVERAGEIF({pnl},">0"),0)', USD, link)
        put(summary, row, 13, f'=IFERROR(AVERAGEIF({pnl},"<0"),0)', USD, link)
        put(summary, row, 14, f"=MAX({pnl})", USD, link)
        put(summary, row, 15, f"=MIN({pnl})", USD, link)
        put(summary, row, 16, f"=MAX({s}$R${a}:$R${b})", "0", link)
        put(summary, row, 17, f"=MAX({s}$Q${a}:$Q${b})", USD, link)
        row += 1
    last_inst = row - 1
    put(summary, row, 1, "Всего", font=bold)
    put(summary, row, 2, "сумма по 8 инструментам, по 1 контракту каждого", font=muted)
    for col, fmt in ((3, "#,##0"), (4, "#,##0"), (5, "#,##0"), (7, USD), (8, USD), (9, USD)):
        letter = get_column_letter(col)
        put(summary, row, col, f"=SUM({letter}5:{letter}{last_inst})", fmt, bold)
    put(summary, row, 6, f"=IF(C{row}>0,D{row}/C{row},0)", "0.0%", bold)
    put(summary, row, 10, f"=IF(C{row}>0,G{row}/C{row},0)", USD, bold)
    for col in range(1, 18):
        summary.cell(row, col).border = Border(top=Side(style="thin", color="7F7F7F"))
    put(summary, row + 2, 1, "WTI (CL) торгуется на MOEX с июля 2026 года (2 181 пятиминутный бар) — справочно; в "
        "выводах исследования не участвует. У COPPER (HG) два года истории.", font=muted)
    put(summary, row + 3, 1, "Листы инструментов — каждая сделка по порядку с накопленным итогом, просадкой и "
        "сериями убытков; «Все сделки» — все инструменты вместе в порядке закрытия.", font=muted)
    summary.freeze_panes = "C5"
    wb.calculation.fullCalcOnLoad = True  # spreadsheet apps compute every formula when the file opens
    wb.save(path)
    return path


def load_ledger(path: Path = LEDGER_CSV) -> pd.DataFrame:
    return pd.read_csv(path, parse_dates=["entry_time", "exit_time"])


def summary_table(led: pd.DataFrame) -> pd.DataFrame:
    """Per-instrument counts and money, as in the workbook's summary sheet."""
    rows = {}
    for name, g in led.groupby("instrument", sort=False):
        p, cum = g["pnl_usd"], g["cum_pnl_usd"].to_numpy()
        losing = (p < 0).to_numpy()
        runs = np.diff(np.flatnonzero(np.r_[True, ~losing, True])) - 1
        rows[name] = {"сделок": len(g), "прибыльных": int((p > 0).sum()), "убыточных": int((p < 0).sum()),
                      "доля прибыльных": (p > 0).mean(), "итого PnL, $": p.sum(),
                      "PnL IS, $": p[g["sample"] == "IS"].sum(), "PnL OOS, $": p[g["sample"] == "OOS"].sum(),
                      "средняя сделка, $": p.mean(), "средняя сделка, тиков": g["pnl_ticks"].mean(),
                      "лучшая, $": p.max(), "худшая, $": p.min(), "макс. убытков подряд": int(runs.max()),
                      "макс. просадка, $": float((np.maximum.accumulate(np.maximum(cum, 0)) - cum).max())}
    return pd.DataFrame(rows).T


def outcomes_figure(led: pd.DataFrame, path: Path) -> Path:
    """Per instrument: every trade's PnL in order (bars) and the running total, IS/OOS boundary marked."""
    from src.ma_slope import plots as P

    plt = P.plt
    names = list(dict.fromkeys(led["instrument"]))
    fig = plt.figure(figsize=(16, 9.5))
    outer = fig.add_gridspec(2, 4, hspace=0.42, wspace=0.28, top=0.93)
    for k, name in enumerate(names):
        g = led[led["instrument"] == name]
        n = np.arange(1, len(g) + 1)
        inner = outer[k // 4, k % 4].subgridspec(2, 1, height_ratios=[2, 1], hspace=0.08)
        top, bottom = fig.add_subplot(inner[0]), fig.add_subplot(inner[1])
        top.plot(n, g["cum_pnl_usd"].to_numpy(), color=P.SERIES[0], lw=1.4)
        pnl = g["pnl_usd"].to_numpy()
        bottom.vlines(n, 0, pnl, colors=np.where(pnl >= 0, P.POS, P.NEG), lw=0.6)
        oos = np.flatnonzero(g["sample"].to_numpy() == "OOS")
        for ax in (top, bottom):
            ax.axhline(0, color=P.BASE, lw=0.8)
            ax.set_xlim(0, len(g) + 1)
            if oos.size:
                ax.axvline(oos[0] + 0.5, color=P.INK2, lw=1)
        top.set_title(f"{name}: {len(g)} trades, total {pnl.sum():+,.0f} $")
        top.tick_params(labelbottom=False)
        bottom.set_xlabel("trade number")
        if k % 4 == 0:
            top.set_ylabel("running total, $")
            bottom.set_ylabel("trade PnL, $")
    fig.suptitle("Every trade of each instrument's selected config in order: PnL per trade (blue gain, red loss) "
                 "and running total, 1 contract; line = IS/OOS boundary", fontsize=10, color=P.INK2, y=0.995)
    return P._save(fig, path)


DECIMALS = {"ng": 3, "br": 2, "gold": 1, "silv": 2, "plt": 1, "pld": 2, "copper": 0, "wti": 2}
OUTCOME_RU = {"profit": "прибыль", "loss": "убыток", "zero": "ноль"}


def format_trades(led: pd.DataFrame, ticker: str) -> str:
    """Every trade of one instrument as a plain-text table (renders compactly in a notebook on GitHub)."""
    g = led[led["ticker"] == ticker]
    d = DECIMALS[ticker]
    out = pd.DataFrame({
        "№": g["trade"].to_numpy(),
        "выборка": g["sample"].to_numpy(),
        "сторона": g["direction"].to_numpy(),
        "вход": g["entry_time"].dt.strftime("%Y-%m-%d %H:%M").to_numpy(),
        "контракт": g["entry_contract"].to_numpy(),
        "цена входа": g["entry_price"].map(lambda v: f"{v:,.{d}f}").to_numpy(),
        "выход": g["exit_time"].dt.strftime("%Y-%m-%d %H:%M").to_numpy(),
        "контракт ": g["exit_contract"].to_numpy(),
        "цена выхода": g["exit_price"].map(lambda v: f"{v:,.{d}f}").to_numpy(),
        "часов": g["hours"].map(lambda v: f"{v:.1f}").to_numpy(),
        "PnL, $": g["pnl_usd"].map(lambda v: f"{v:+,.2f}").to_numpy(),
        "тиков": g["pnl_ticks"].map(lambda v: f"{v:+,.1f}").to_numpy(),
        "исход": g["outcome"].map(OUTCOME_RU).to_numpy(),
        "накопленный, $": g["cum_pnl_usd"].map(lambda v: f"{v:+,.2f}").to_numpy(),
    })
    return out.to_string(index=False)
