"""Проверка журналов сделок: все ли сделки оптимизированных систем на месте и по порядку.

Каждый журнал пересчитывается заново способом, не зависящим от того, как он выгружен:

* скользящие средние — событийным ядром исследования (kernel.simulate). Финалисты
  выгружались из VectorBT, поэтому для них это независимая сверка: совпасть должна
  каждая сделка — даты входа и выхода, направление, котировки, результат. Число
  сделок и PnL IS/OOS — с полной сеткой, шаги walk-forward — с wf_<критерий>.csv;
* семь семейств — повторным прогоном движка пакета; число сделок OOS и итог OOS — с
  отчётом, число сделок IS — с results/sweeps.csv, walk-forward — с разделом 3;
* структура каждого журнала: нумерация 1..n без пропусков, порядок по дате выхода,
  вход не позже выхода, у систем с одной позицией следующая сделка открывается
  после закрытия предыдущей, нет повторов, период соответствует дате выхода, без
  ролла результат равен разности котировок;
* производные файлы: листы Excel, таблицы в пунктах и столбцы «только числа»
  повторяют журнал строка в строку.

    python audit_trades.py ma      # окружение исследования (numba, pandas<3): скользящие средние
    python audit_trades.py seven   # окружение пакета exchange: семь семейств, таблицы в пунктах
    python audit_trades.py files   # окружение с openpyxl: книги Excel; затем AUDIT.md

Результаты шагов — audit/ma.json, audit/seven.json, audit/files.json; сводка —
AUDIT.md и audit/systems.csv (по строке на систему). Ненулевой код возврата — есть
расхождение.
"""
from __future__ import annotations

import json
import sys
import warnings
from pathlib import Path

import numpy as np
import pandas as pd

HERE = Path(__file__).resolve().parent
CSV = HERE / "csv"
PTS = HERE / "points"
OUT = HERE / "audit"
ROOT = HERE.parents[1]
MA_RESULTS = ROOT / "research" / "ma_crossover_5m" / "results"
OOS_START = pd.Timestamp("2024-01-01")
MULT = {"NG": 100.0, "BR": 10.0}
# Даты в журналах двух видов (день и минута); разбор поэлементно корректен, предупреждение лишнее.
warnings.filterwarnings("ignore", message="Could not infer format")


def read(path: Path) -> pd.DataFrame:
    return pd.read_csv(path, keep_default_na=False, na_values=[""])


def structure(df: pd.DataFrame, *, single: bool, period_rule=None, market: str | None = None,
              ordered_by=("выход",)) -> list[str]:
    """Проверки формы журнала; возвращает список нарушений (пустой — всё в порядке)."""
    bad = []
    if not np.array_equal(df["№"].to_numpy(), np.arange(1, len(df) + 1)):
        bad.append("нумерация не 1..n")
    trades = df[df["вход"].notna()]
    ent, ext = pd.to_datetime(trades["вход"]), pd.to_datetime(trades["выход"])
    if (ent > ext).any():
        bad.append(f"вход позже выхода: {(ent > ext).sum()}")
    keys = [c for c in ordered_by if c != "выход"]
    for _, g in (df.groupby(keys, sort=False) if keys else [(None, df)]):
        x = pd.to_datetime(g["выход"])
        if not x.is_monotonic_increasing:
            bad.append("сделки не по порядку закрытия")
            break
    if single and len(trades) > 1:
        overlap = ent.to_numpy()[1:] <= ext.to_numpy()[:-1]
        if overlap.any():
            bad.append(f"следующая сделка открыта до закрытия предыдущей: {overlap.sum()}")
    dup = trades.duplicated(["нога", "вход", "выход", "направление", "результат, USD"]).sum()
    if dup:
        bad.append(f"повторяющиеся строки: {dup}")
    if period_rule is not None:
        expected = period_rule(pd.to_datetime(df["выход"]))
        if (np.asarray(expected) != df["период"].to_numpy()).any():
            bad.append("период не соответствует дате выхода")
    if market is not None:                                    # 1 контракт: без ролла — разность котировок
        plain = trades[trades["через ролл"].fillna("") == ""]
        sign = np.where(plain["направление"] == "long", 1.0, -1.0)
        quoted = (plain["котировка выхода"] - plain["котировка входа"]) * sign * MULT[market]
        miss = (quoted - plain["результат, USD"]).abs() > 0.011
        if miss.any():
            bad.append(f"результат без ролла не равен разности котировок: {miss.sum()}")
    if df["результат, USD"].isna().any():
        bad.append("пустой результат")
    return bad


def same_frame(a: pd.DataFrame, b: pd.DataFrame, columns) -> list[str]:
    """Построчное сравнение двух журналов по колонкам (числа — до копейки)."""
    if len(a) != len(b):
        return [f"строк {len(a)} против {len(b)}"]
    bad = []
    for c in columns:
        x, y = a[c].reset_index(drop=True), b[c].reset_index(drop=True)
        if c in ("вход", "выход"):
            ok = (pd.to_datetime(x).fillna(pd.Timestamp(0)) == pd.to_datetime(y).fillna(pd.Timestamp(0))).all()
        elif pd.api.types.is_numeric_dtype(x) and pd.api.types.is_numeric_dtype(y):
            ok = np.allclose(x.to_numpy(float), y.to_numpy(float), rtol=1e-9, atol=0.006, equal_nan=True)
        else:
            ok = (x.fillna("").astype(str) == y.fillna("").astype(str)).all()
        if not ok:
            bad.append(f"не совпала колонка «{c}»")
    return bad


def record(kind, sid, path, df, issues, **extra) -> dict:
    trades = df[df["вход"].notna()]
    return {"kind": kind, "id": sid, "file": str(path.relative_to(HERE)), "rows": len(df), "trades": len(trades),
            "first_entry": str(pd.to_datetime(trades["вход"]).min()) if len(trades) else "",
            "last_exit": str(pd.to_datetime(df["выход"]).max()) if len(df) else "",
            "pnl": round(float(df["результат, USD"].sum()), 2), "issues": issues, **extra}


# ===================================================================== скользящие средние
def audit_ma() -> list[dict]:
    import ma_kernel as MK
    from export_trades import WF_COLUMNS, finalist_id

    def ma_period(x):
        return np.where(x < OOS_START, "IS", "OOS")

    out = []
    families = {m["id"]: m for m in json.loads((CSV / "ma_families_meta.json").read_text())}
    wf_meta = {m["id"]: m for m in json.loads((CSV / "ma_wf_meta.json").read_text())}
    for stage in ("stage1", "stage3"):
        for t in ("NG", "BR"):
            kern, slope = MK.Kernel(t), MK.slopes(stage, t)
            # --- финалисты: VectorBT-журнал против ядра, число сделок — против сетки
            finalists = json.loads((MA_RESULTS / stage / t / "finalists.json").read_text())
            gcols = ["config", "trades_is", "trades_oos", "pnl_is", "pnl_oos"]
            g = pd.concat([pd.read_parquet(MA_RESULTS / stage / t / f"grid_{m}.parquet", columns=gcols,
                                           filters=[("config", "in", [int(f["grid_config"]) for f in finalists])])
                           for m in ("A_pct", "B_atr")]).set_index("config")
            for f in finalists:
                sid = finalist_id(stage, t, f)
                path = CSV / (f"{sid}.csv" if f["scope"] == "global" else f"groups/{sid}.csv.gz")
                df = read(path)
                ref = kern.journal(f["config"], slope[f["config"]["method"]])
                issues = structure(df, single=True, period_rule=ma_period, market=t)
                issues += [f"ядро: {x}" for x in same_frame(df, ref, MK.COLUMNS)]
                row = g.loc[int(f["grid_config"])]
                for p in ("IS", "OOS"):
                    part = df[df["период"] == p]
                    if len(part) != int(row[f"trades_{p.lower()}"]):
                        issues.append(f"сделок {p} {len(part)} против сетки {int(row[f'trades_{p.lower()}'])}")
                    if abs(part["результат, USD"].sum() - row[f"pnl_{p.lower()}"]) > 0.51:
                        issues.append(f"PnL {p} против сетки")
                if abs(df["результат, USD"].sum() - f["full"]["pnl"]) > 0.05:
                    issues.append("итог против VectorBT-статистики финалиста")
                out.append(record("ma_finalist", sid, path, df, issues, stage=stage, market=t,
                                  scope=f["scope"], rule=f["rule"], config=int(f["grid_config"]),
                                  reference="VectorBT-журнал = событийное ядро; сделки IS/OOS = сетка"))
            print(f"{stage} {t}: финалистов {len(finalists)}", flush=True)
            # --- все правила отбора указывают на журнал (совпавшие конфигурации записаны один раз)
            picks = pd.read_csv(MA_RESULTS / stage / t / "picks.csv")
            ids = {int(f["grid_config"]): finalist_id(stage, t, f) for f in finalists}
            for p in picks[~picks["scope"].str.startswith("family:")].itertuples():
                own = finalist_id(stage, t, {"scope": p.scope, "rule": p.rule,
                                             "config": {"exit": p.exit, "method": p.method}})
                target = ids.get(int(p.config))
                out.append({"kind": "ma_pick", "id": own, "stage": stage, "market": t, "scope": p.scope,
                            "rule": p.rule, "config": int(p.config), "file_id": target or "",
                            "issues": [] if target else ["нет журнала для этого отбора"]})
            # --- семейства: файл против повторного пересчёта и против сетки
            fam = picks[picks["scope"].str.startswith("family:")]
            for p in fam.itertuples(index=False):
                cfg = {k: getattr(p, k) for k in MK.PARAMS}
                sid = (f"MA_{stage[-1]}_{t}_f_{cfg['fast_type']}-{cfg['slow_type']}"
                       f"_x{cfg['exit'][0]}_s{cfg['method'][0]}")
                path = CSV / "families" / f"{sid}.csv.gz"
                if not path.exists() or sid not in families:
                    out.append({"kind": "ma_family", "id": sid, "file": str(path.relative_to(HERE)),
                                "issues": ["нет файла семейства"]})
                    continue
                df = read(path)
                ref = kern.journal(cfg, slope[cfg["method"]])
                issues = structure(df, single=True, period_rule=ma_period, market=t)
                issues += [f"пересчёт: {x}" for x in same_frame(df, ref, MK.COLUMNS)]
                for per, n_ref, pnl_ref in (("IS", p.trades_is, p.pnl_is), ("OOS", p.trades_oos, p.pnl_oos)):
                    part = df[df["период"] == per]
                    if len(part) != int(n_ref) or abs(part["результат, USD"].sum() - pnl_ref) > 0.51:
                        issues.append(f"{per} против сетки")
                if families[sid]["grid_config"] != int(p.config):
                    issues.append("номер конфигурации не совпал с picks.csv")
                same = families[sid]["same_as"]
                if same:                                       # совпавшая с финалистом — те же сделки
                    fpath = CSV / (f"{same}.csv" if "_x" not in same else f"groups/{same}.csv.gz")
                    issues += [f"против {same}: {x}" for x in same_frame(df, read(fpath), MK.COLUMNS)]
                out.append(record("ma_family", sid, path, df, issues, stage=stage, market=t,
                                  scope=p.scope, rule="R3_plateau", config=int(p.config), same_as=same,
                                  reference="пересчёт ядром; сделки и PnL IS/OOS = сетка (picks.csv)"))
            print(f"{stage} {t}: семейств {len(fam)}", flush=True)
            # --- walk-forward: каждый шаг против ядра и против wf_<критерий>.csv
            months = kern.mk.months
            for crit in ("pnl", "sharpe", "plateau"):
                sid = f"MA_{stage[-1]}_{t}_WF_{crit}"
                path = CSV / "walkforward" / f"{sid}.csv.gz"
                df = read(path)
                wf = pd.read_csv(MA_RESULTS / stage / t / f"wf_{crit}.csv")
                params = MK.config_params(stage, t, wf["config"])
                issues = structure(df, single=False, period_rule=ma_period, market=t, ordered_by=("шаг", "выход"))
                if list(df.columns) != WF_COLUMNS:
                    issues.append("колонки не те")
                if sorted(df["шаг"].unique()) != list(range(1, len(wf) + 1)):
                    issues.append("не все шаги на месте")
                for step, r in enumerate(wf.itertuples(index=False), 1):
                    part = df[df["шаг"] == step]
                    cfg = params[int(r.config)]
                    ref, month = kern.journal_months(cfg, slope[cfg["method"]])
                    ref = ref[(month >= r.test_start) & (month < r.test_end)]
                    cols = [c for c in MK.COLUMNS if c not in ("№", "примечание")]
                    issues += [f"шаг {step}: {x}" for x in same_frame(part, ref, cols)]
                    if abs(part["результат, USD"].sum() - r.test_pnl) > 0.15:
                        issues.append(f"шаг {step}: итог против wf_{crit}.csv")
                    if (part["номер в сетке"] != int(r.config)).any():
                        issues.append(f"шаг {step}: не та конфигурация")
                    inside = pd.to_datetime(part["выход"]).dt.to_period("M")
                    if ((inside < months[r.test_start]) | (inside > months[r.test_end - 1])).any():
                        issues.append(f"шаг {step}: выход вне месяцев шага")
                out.append(record("ma_wf", sid, path, df, issues, stage=stage, market=t, rule=crit,
                                  steps=len(wf), reference=f"пересчёт ядром по шагам; итог шага = wf_{crit}.csv",
                                  entered_before_step=wf_meta[sid]["trades_entered_before_step"]))
            del kern
    return out


# ===================================================================== семь семейств
def audit_seven() -> list[dict]:
    from export_trades import COLUMNS, SEVEN_DOC, WF7_COLUMNS, Seven, dates, md_table, seven_cells

    from exchange.stats.cv import WalkForward

    S = Seven()
    sweeps = pd.read_csv(ROOT / "research" / "seven_families" / "results" / "sweeps.csv")
    winners = sweeps[sweeps["is_winner"]].set_index(["family", "market"])
    wf_ref = md_table(SEVEN_DOC.read_text(), "WF эффект.").set_index("ячейка")
    out = []
    cols = [c for c in COLUMNS if c != "№"]

    def frame(rows, order):
        if not rows:                                   # окно без сделок
            return pd.DataFrame(columns=COLUMNS)
        df = dates(pd.DataFrame(rows).sort_values(order, kind="stable").reset_index(drop=True))
        df.insert(0, "№", np.arange(1, len(df) + 1))
        return df

    for cell in seven_cells().itertuples():
        family, market, params = cell.family, cell.market, cell.params
        index = S.index(family, market)
        is_idx, oos_idx = S.RA.split(index)
        # --- журнал отчёта: два прогона
        path = CSV / f"{cell.id}.csv"
        df = read(path)
        rows, totals = [], {}
        for period, idx in (("IS", is_idx), ("OOS", oos_idx)):
            r, totals[period] = S.run(family, market, params, idx, period)
            rows += r
        ref = frame(rows, ["период", "выход", "нога"])
        issues = structure(df, single=False, ordered_by=("период", "выход"))
        issues += [f"повторный прогон: {x}" for x in same_frame(df, ref, cols)]
        n_oos = int(((df["период"] == "OOS") & df["вход"].notna()).sum())
        n_is = int(((df["период"] == "IS") & df["вход"].notna()).sum())
        if n_oos != int(cell.сделок):
            issues.append(f"сделок OOS {n_oos} против отчёта {cell.сделок}")
        w = winners.loc[(family, market)]
        if n_is != int(w["is_trades"]):
            issues.append(f"сделок IS {n_is} против sweeps.csv {int(w['is_trades'])}")
        if abs(df.loc[df["период"] == "OOS", "результат, USD"].sum() / 1e6 - cell.oos_return_report) > 0.0006:
            issues.append("итог OOS против отчёта")
        is_end, oos_start = is_idx[-1], oos_idx[0]
        first_oos = pd.to_datetime(df.loc[(df["период"] == "OOS") & df["вход"].notna(), "вход"]).min()
        out.append(record("seven_report", cell.id, path, df, issues, market=market, params=cell.лучшая,
                          is_end=str(is_end.date()), oos_start=str(oos_start.date()),
                          first_oos_entry=str(first_oos.date()) if pd.notna(first_oos) else "",
                          reference="повторный прогон движка; сделок OOS = отчёт, сделок IS = sweeps.csv"))
        # --- сквозной прогон
        path = CSV / "full" / f"{cell.id}.csv"
        full = read(path)

        def period(ts, is_end=is_end, oos_start=oos_start):
            return "IS" if ts <= is_end else ("эмбарго" if ts < oos_start else "OOS")

        rows, change = S.run(family, market, params, index, period)
        ref = frame(rows, ["выход", "нога"])
        issues = structure(full, single=False,
                           period_rule=lambda x, is_end=is_end, oos_start=oos_start: np.where(
                               x <= is_end.tz_localize(None), "IS",
                               np.where(x < oos_start.tz_localize(None), "эмбарго", "OOS")))
        issues += [f"повторный прогон: {x}" for x in same_frame(full, ref, cols)]
        if abs(full["результат, USD"].sum() - change) > 0.01:
            issues.append("сумма сделок не равна изменению капитала")
        a = full[(full["период"] == "IS") & full["вход"].notna()]
        b = df[(df["период"] == "IS") & df["вход"].notna()]
        if same_frame(a.sort_values(["выход", "нога", "вход"], kind="stable"),
                      b.sort_values(["выход", "нога", "вход"], kind="stable"), ["нога", "вход", "выход", "результат, USD"]):
            issues.append("сделки до конца IS не совпали с журналом отчёта")
        span = full[pd.to_datetime(full["выход"]) > is_end.tz_localize(None)]
        gap = span[pd.to_datetime(span["выход"]) < first_oos] if pd.notna(first_oos) else span
        out.append(record("seven_full", cell.id, path, full, issues, market=market, params=cell.лучшая,
                          is_end=str(is_end.date()), oos_start=str(oos_start.date()),
                          gap_trades=int(gap["вход"].notna().sum()), gap_pnl=round(float(gap["результат, USD"].sum()), 2),
                          reference="повторный прогон; сумма = изменение капитала; до конца IS = журнал отчёта"))
        # --- walk-forward: сделки окна против прогона с параметрами окна
        path = CSV / "walkforward" / f"{cell.id}_WF.csv"
        wf = read(path)
        meta = next(m for m in json.loads((CSV / "seven_wf_meta.json").read_text()) if m["id"] == f"{cell.id}_WF")
        issues = structure(wf, single=False, ordered_by=("окно", "выход"))
        if list(wf.columns) != WF7_COLUMNS:
            issues.append("колонки не те")
        split = WalkForward(train_size=max(int(len(index) * 0.35), 250),
                            test_size=max(int(len(index) * 0.08), 120), horizon=5, embargo=S.RA.EMBARGO)
        tests = [index[sp.test] for sp in split.split(len(index))]
        for win in meta["windows"]:
            part = wf[wf["окно"] == win["окно"]]
            idx = tests[win["окно"] - 1]
            if win["тест"] != f"{idx[0].date()}…{idx[-1].date()}":
                issues.append(f"окно {win['окно']}: даты теста не те")
            p = {k: (float(v) if "." in v else int(v)) for k, v in (x.split("=") for x in win["параметры"].split(","))}
            rows, change = S.run(family, market, p, idx, "WF")
            ref = frame(rows, ["выход", "нога"])
            issues += [f"окно {win['окно']}: {x}" for x in same_frame(part, ref, cols)]
            if abs(part["результат, USD"].sum() - change) > 0.01:
                issues.append(f"окно {win['окно']}: сумма не равна изменению капитала")
        rep = wf_ref.loc[cell.ячейка]
        for k, v in meta["report"].items():
            if rep[k] != "" and abs(v - float(rep[k])) > 0.0006:
                issues.append(f"{k} против отчёта")
        out.append(record("seven_wf", f"{cell.id}_WF", path, wf, issues, market=market,
                          windows=len(meta["windows"]),
                          reference="повторный прогон окон; эффективность, доля окон>0, стабильность = раздел 3"))
        print(f"{cell.id:22} проверен", flush=True)
    return out


def audit_points() -> list[dict]:
    """Таблицы в пунктах и столбцы «только числа» — против журналов сделок."""
    out = []
    for mode, pdir, cdir in (("report", PTS, CSV), ("full", PTS / "full", CSV / "full")):
        seq = PTS / "sequence" / ("" if mode == "report" else "full")
        meta = json.loads((pdir / "seven_meta.json").read_text())
        if mode == "report":
            meta += json.loads((pdir / "ma_meta.json").read_text())
        for m in meta:
            path = pdir / f"{m['id']}.csv"
            pts = read(path)
            # у портфельных стратегий таблица на ногу (…_NGBR_NG), журнал — общий (…_NGBR)
            src_id = m["id"] if (cdir / f"{m['id']}.csv").exists() else m["id"].rsplit("_", 1)[0]
            log = read(cdir / f"{src_id}.csv")
            log = log[log["нога"] == m["asset"]].reset_index(drop=True)
            issues = []
            if len(pts) != len(log):
                issues.append(f"строк {len(pts)} против журнала {len(log)}")
            else:
                for a, b in (("вход", "вход"), ("выход", "выход")):
                    x = pd.to_datetime(pts[a]).fillna(pd.Timestamp(0))
                    y = pd.to_datetime(log[b]).fillna(pd.Timestamp(0))
                    if (x != y).any():
                        issues.append(f"не совпала колонка «{a}»")
                if not np.allclose(pts["результат, USD"], log["результат, USD"], atol=0.006):
                    issues.append("результат в USD не совпал с журналом")
                if (pts["период"].astype(str) != log["период"].astype(str)).any():
                    issues.append("период не совпал с журналом")
            steps = pts.loc[pts["вход"].notna(), "результат, шагов цены"].to_numpy(float)
            whole = (np.sign(steps) * np.floor(np.abs(steps) + 0.5)).astype(int)
            seq_vals = np.loadtxt(seq / f"{m['id']}.csv", dtype=int, ndmin=1)
            if not np.array_equal(seq_vals, whole):
                issues.append("файл «только числа» не совпал с таблицей в пунктах")
            out.append({"kind": f"points_{mode}", "id": m["id"], "file": str(path.relative_to(HERE)),
                        "rows": len(pts), "trades": int(pts["вход"].notna().sum()), "sequence": len(seq_vals),
                        "issues": issues})
    return out


# ===================================================================== книги Excel и сводка
def audit_files() -> list[dict]:
    from openpyxl import load_workbook

    import build_excel as B

    out = []
    for kind, book in (("main", B.OUT), ("groups", B.OUT_GROUPS), ("full", B.OUT_FULL)):
        wb = load_workbook(book, read_only=True)
        for m in B.strategies(kind):
            df = read(CSV / m["csv"])
            ws = wb[m["id"]]
            got = list(ws.iter_rows(min_row=B.FIRST, max_col=12, values_only=True))
            got = [r for r in got if r[0] is not None]
            issues = []
            if len(got) != len(df):
                issues.append(f"строк на листе {len(got)} против CSV {len(df)}")
            else:
                num = np.array([r[0] for r in got])
                res = np.array([r[11] for r in got], dtype=float)
                exits = pd.to_datetime(pd.Series([r[4] for r in got]))
                if not np.array_equal(num, df["№"].to_numpy()):
                    issues.append("номера сделок")
                if not np.allclose(res, df["результат, USD"].to_numpy(float), atol=0.006):
                    issues.append("результаты сделок")
                if (exits.to_numpy() != pd.to_datetime(df["выход"]).to_numpy()).any():
                    issues.append("даты выхода")
            out.append({"kind": f"excel_{kind}", "id": m["id"], "file": book.name, "rows": len(got), "issues": issues})
        wb.close()
    for book, pdir in ((PTS / "points.xlsx", PTS), (PTS / "full" / "points_full.xlsx", PTS / "full")):
        wb = load_workbook(book, read_only=True)
        meta = json.loads((pdir / "seven_meta.json").read_text())
        if (pdir / "ma_meta.json").exists():
            meta += json.loads((pdir / "ma_meta.json").read_text())
        for m in meta:
            df = read(pdir / f"{m['id']}.csv")
            ws = wb[m["id"][:31]]
            head = next(ws.iter_rows(min_row=B.FIRST - 1, max_row=B.FIRST - 1, values_only=True))
            col = list(head).index("результат, USD")
            got = [r for r in ws.iter_rows(min_row=B.FIRST, values_only=True) if r[0] is not None]
            issues = []
            if len(got) != len(df):
                issues.append(f"строк на листе {len(got)} против CSV {len(df)}")
            elif not np.allclose(np.array([r[col] for r in got], dtype=float), df["результат, USD"], atol=0.006):
                issues.append("результаты сделок")
            out.append({"kind": "excel_points", "id": m["id"], "file": str(book.relative_to(HERE)), "rows": len(got),
                        "issues": issues})
        wb.close()
    return out


def load(name: str) -> list[dict]:
    path = OUT / f"{name}.json"
    if not path.exists():
        raise SystemExit(f"нет {path}: сначала python audit_trades.py {name}")
    return json.loads(path.read_text())


KIND = {
    "seven_report": "Семь семейств: прогоны отчёта (IS и OOS отдельно)",
    "seven_full": "Семь семейств: сквозной прогон по всей истории",
    "seven_wf": "Семь семейств: walk-forward, окна",
    "ma_finalist": "Скользящие средние: глобальные отборы и лучшие в группах",
    "ma_family": "Скользящие средние: лучший в каждом из 144 семейств",
    "ma_wf": "Скользящие средние: walk-forward, шаги",
}


REFERENCE = {
    "seven_report": "повторный прогон движка; сделок OOS — с отчётом, сделок IS — с `sweeps.csv`, итог OOS — с отчётом",
    "seven_full": "повторный прогон; сумма сделок = изменение капитала; сделки до конца IS = журнал отчёта",
    "seven_wf": "повторный прогон каждого окна; эффективность, доля окон > 0 и стабильность параметров — с разделом 3",
    "ma_finalist": "VectorBT-журнал против событийного ядра сделка в сделку; сделки и PnL IS/OOS — с сеткой",
    "ma_family": "пересчёт ядром; сделки и PnL IS/OOS — с сеткой (`picks.csv`)",
    "ma_wf": "пересчёт ядром по шагам; итог шага — с `wf_<критерий>.csv`",
}


def sp(v) -> str:
    return f"{int(v):,}".replace(",", " ")


def money(v) -> str:
    return f"{v:+,.2f}".replace(",", " ")


def write_report(ma, seven, files):
    systems = pd.DataFrame([r for r in ma + seven if r["kind"] in KIND])
    systems["ok"] = systems["issues"].map(len) == 0
    picks = pd.DataFrame([r for r in ma if r["kind"] == "ma_pick"])
    points = pd.DataFrame([r for r in seven if r["kind"].startswith("points")])
    excel = pd.DataFrame(files)
    problems = [(r["kind"], r["id"], x) for r in ma + seven + files for x in r["issues"]]
    OUT.mkdir(exist_ok=True)
    cols = ["kind", "id", "stage", "market", "scope", "rule", "config", "file", "rows", "trades", "first_entry",
            "last_exit", "pnl", "ok", "issues", "reference"]
    flat = systems.reindex(columns=cols).assign(issues=systems["issues"].map("; ".join))
    flat.to_csv(OUT / "systems.csv", index=False)

    lines = ["# Проверка журналов сделок", "",
             "Вопрос: есть ли **все** сделки **каждой** системы, которую отобрала оптимизация, и идут ли они по "
             "порядку. Ответ даёт `audit_trades.py`: каждый журнал пересчитан заново способом, не зависящим от "
             "того, как он был выгружен, и сверен построчно.", ""]
    lines += ["## Итог", ""]
    total_trades = systems["trades"].fillna(0).sum()
    if problems:
        lines += [f"**Найдено расхождений: {len(problems)}.** Список — в конце файла.", ""]
    else:
        lines += [f"**Расхождений нет.** Систем: {len(systems)}, сделок: {sp(total_trades)}. Каждая сделка "
                  "каждой системы на месте, по порядку, без пропусков и повторов; производные файлы (Excel, "
                  "пункты, «только числа») повторяют журналы строка в строку.", ""]
    lines += ["| Что проверено | Систем | Сделок | Без расхождений | Против чего сверено |", "|---|---|---|---|---|"]
    for kind, title in KIND.items():
        g = systems[systems["kind"] == kind]
        lines.append(f"| {title} | {len(g)} | {sp(g['trades'].sum())} | {int(g['ok'].sum())} из {len(g)} "
                     f"| {REFERENCE[kind]} |")
    lines.append(f"| Таблицы в пунктах и «только числа» | {len(points)} | {sp(points['trades'].sum())} "
                 f"| {int((points['issues'].map(len) == 0).sum())} из {len(points)} | журналы сделок строка в строку |")
    lines.append(f"| Листы Excel | {len(excel)} | {sp(excel['rows'].sum())} строк "
                 f"| {int((excel['issues'].map(len) == 0).sum())} из {len(excel)} | CSV строка в строку |")
    lines += ["", "Проверки каждого журнала:", "",
              "* **полнота** — повторный расчёт даёт ровно те же сделки: у скользящих средних — событийное ядро "
              "исследования против файлов (финалисты выгружены из VectorBT, так что это сверка двух независимых "
              "реализаций), у семи семейств — повторный прогон движка; число сделок IS и OOS — с полной сеткой "
              "(скользящие средние), с отчётом и `sweeps.csv` (семь семейств);",
              "* **порядок** — нумерация 1..n без пропусков, сделки по дате закрытия, вход не позже выхода; у систем "
              "с одной позицией (скользящие средние) следующая сделка открывается только после закрытия предыдущей;",
              "* **повторы** — ни одна сделка не записана дважды;",
              "* **разметка** — период IS/OOS соответствует дате выхода; без ролла результат 1 контракта равен "
              "разности котировок × множитель.", ""]

    lines += ["## Что было неполным и что добавлено", "",
              "1. **Семь семейств: разрыв между IS и OOS.** Отчёт оценивает IS и OOS отдельными прогонами: OOS "
              "начинается без позиции и с пустыми индикаторами, между периодами — эмбарго 10 дней. Поэтому в "
              "журнале отчёта нет сделок за несколько месяцев после конца IS, а открытая позиция на конец IS "
              "закрыта переоценкой. Добавлен **сквозной прогон** той же системы по всей истории "
              "([`csv/full/`](csv/full/), [`trade_log_full.xlsx`](trade_log_full.xlsx), "
              "[`points/full/`](points/full/README.md)): в нём сделки идут подряд с первого до последнего дня данных. "
              "Его сделки до конца IS совпадают с журналом отчёта одна в одну.", ""]
    rep = systems[systems["kind"] == "seven_report"].set_index("id")
    full = systems[systems["kind"] == "seven_full"].set_index("id")
    lines += ["| Система | Конец IS | Первая сделка OOS в прогоне отчёта | Сделок сквозного прогона в этом окне "
              "| Их итог, USD |", "|---|---|---|---|---|"]
    for sid, r in rep.iterrows():
        f = full.loc[sid]
        lines.append(f"| {sid} | {r['is_end']} | {r['first_oos_entry']} | {int(f['gap_trades'])} "
                     f"| {money(f['gap_pnl'])} |")
    lines += ["",
              "2. **Скользящие средние: 144 семейства.** Разделы 7 и 12 отчёта ранжируют лучшую по плато конфигурацию "
              "каждого семейства (сочетание типов × выход × метод), но их сделок не было. Добавлены все: "
              f"[`csv/families/`](csv/families/README.md), {len(systems[systems['kind'] == 'ma_family'])} файлов "
              "на этапах 1 и 3.",
              "3. **Walk-forward с переоптимизацией** — тоже оптимизированная система: на каждом шаге свои "
              "параметры. Добавлены сделки всех шагов: [`csv/walkforward/`](csv/walkforward/), 12 систем семи "
              "семейств и 12 — скользящих средних (2 этапа × 2 рынка × 3 критерия). Итоги шагов совпадают с "
              "отчётом. У скользящих средних сделка относится к шагу по месяцу выхода, как в отчёте; "
              f"вошедших до начала своего шага — {sp(systems.loc[systems['kind'] == 'ma_wf', 'entered_before_step'].sum())}, "
              "они помечены в «примечании».", ""]

    lines += ["## Совпадения: одна конфигурация — один журнал", "",
              "Если два правила отбора выбрали одну и ту же конфигурацию, её сделки записаны один раз. Ниже — каждое "
              "правило, у которого нет собственного файла, и журнал, где лежат его сделки.", "",
              "| Этап | Рынок | Отбор | Правило | Конфигурация | Журнал |", "|---|---|---|---|---|---|"]
    for r in picks[picks["id"] != picks["file_id"]].itertuples():
        group = "по всей сетке" if r.scope == "global" else "группа " + r.scope.replace("|", " × ")
        lines.append(f"| {r.stage[-1]} | {r.market} | {group} | {r.rule} | {r.config} | {r.file_id} |")
    fam = systems[(systems["kind"] == "ma_family") & (systems["same_as"].fillna("") != "")]
    lines += ["", f"Кроме того, {len(fam)} лучших в семействах совпали с финалистами; у них есть и свой файл, "
              "и сделки в нём проверены на равенство журналу финалиста (колонка «Совпадает с» в "
              "[`csv/families/README.md`](csv/families/README.md)).", ""]

    lines += ["## По системам", "",
              "Строка на каждую систему — файл, число сделок, первая и последняя сделка, итог и результат проверки: "
              "[`audit/systems.csv`](audit/systems.csv).", ""]
    lines += ["| Система | Файл | Сделок | Первый вход | Последний выход | Итог, USD | Проверка |",
              "|---|---|---|---|---|---|---|"]
    for r in systems[systems["kind"] != "ma_family"].itertuples():
        lines.append(f"| {r.id} | [{Path(r.file).name}]({r.file}) | {r.trades} | {str(r.first_entry)[:16]} "
                     f"| {str(r.last_exit)[:16]} | {money(r.pnl)} | {'✅' if r.ok else '❌ ' + '; '.join(r.issues)} |")
    lines += ["", "Семейства (576) — в `audit/systems.csv`.", ""]
    if problems:
        lines += ["## Расхождения", "", *[f"* `{k}` {i}: {x}" for k, i, x in problems], ""]
    lines += ["Пересборка: `python audit_trades.py ma` (окружение исследования), `python audit_trades.py seven` "
              "(окружение пакета `exchange`), `python audit_trades.py files`."]
    (HERE / "AUDIT.md").write_text("\n".join(lines) + "\n")
    return problems


def save(name: str, rows: list[dict]):
    OUT.mkdir(exist_ok=True)
    (OUT / f"{name}.json").write_text(json.dumps(rows, ensure_ascii=False, indent=1, default=str))
    bad = [(r["id"], x) for r in rows for x in r["issues"]]
    print(f"{name}: проверено {len(rows)}, расхождений {len(bad)}")
    for b in bad[:30]:
        print("  ", *b)
    return bad


if __name__ == "__main__":
    step = sys.argv[1] if len(sys.argv) > 1 else ""
    if step == "ma":
        bad = save("ma", audit_ma())
    elif step == "seven":
        bad = save("seven", audit_seven() + audit_points())
    elif step == "files":
        bad = save("files", audit_files())
        bad += write_report(load("ma"), load("seven"), load("files"))
        print("AUDIT.md записан")
    else:
        raise SystemExit(__doc__)
    sys.exit(1 if bad else 0)
