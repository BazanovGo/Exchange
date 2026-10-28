"""points/README.md и points/points.xlsx из таблиц исходов в пунктах (export_points.py).

README — статистика по каждой таблице и ссылка на неё. Excel — лист на таблицу:
исход и серия убыточных подряд — формулы; лист «Сводка» считает статистику
формулами по листам.
"""
from __future__ import annotations

import json
from pathlib import Path

import pandas as pd
from build_excel import HEAD_FILL, THIN, USD, font
from openpyxl import Workbook
from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter

HERE = Path(__file__).resolve().parent
PTS = HERE / "points"
FIRST = 6
PT_FMT = '0.00000;[Red]-0.00000;0'
NUM_FMT = {"№": "0", "объём, ед. актива": "#,##0.00", "цена входа": "0.0000", "цена выхода": "0.0000",
           "результат, пунктов": PT_FMT, "результат, шагов цены": '#,##0.00;[Red]-#,##0.00;0',
           "результат, USD": USD, "итог, USD": USD, "итог, пунктов": PT_FMT,
           "при издержках 1 шаг, пунктов": PT_FMT}
WIDTH = {"№": 7, "период": 8, "вход": 17, "выход": 17, "направление": 12, "объём, ед. актива": 13,
         "цена входа": 11, "цена выхода": 11, "через ролл": 9, "результат, пунктов": 13, "исход": 10,
         "серия убыточных подряд": 12, "результат, шагов цены": 13, "результат, USD": 14, "итог, USD": 15,
         "итог, пунктов": 13, "при издержках 1 шаг, пунктов": 15, "примечание": 58}


def tables(pts: Path = PTS) -> list[dict]:
    meta = json.loads((pts / "seven_meta.json").read_text())
    if (pts / "ma_meta.json").exists():
        meta += json.loads((pts / "ma_meta.json").read_text())
    for m in meta:
        m["df"] = pd.read_csv(pts / f"{m['id']}.csv", keep_default_na=False, na_values=[""])
    return meta


def stats(df: pd.DataFrame) -> dict:
    trades = df[df["вход"].notna()]
    pts = trades["результат, пунктов"]
    best = streak = 0
    for v in pts:
        streak = streak + 1 if v < 0 else 0
        best = max(best, streak)
    wins, losses = pts[pts > 0], pts[pts < 0]
    usd = df["результат, USD"]
    return {"n": len(trades), "wins": len(wins), "losses": len(losses), "rate": len(wins) / max(len(trades), 1),
            "avg_win": wins.mean(), "avg_loss": losses.mean(), "best": pts.max(), "worst": pts.min(),
            "streak": best, "usd_is": usd[df["период"] == "IS"].sum(), "usd_oos": usd[df["период"] == "OOS"].sum(),
            "usd": usd.sum(), "points": pts.sum() if "итог, пунктов" in df else None,
            "pf": trades.loc[trades["результат, USD"] > 0, "результат, USD"].sum()
            / max(-trades.loc[trades["результат, USD"] < 0, "результат, USD"].sum(), 1e-9)}


def sheet_columns(df: pd.DataFrame) -> list[str]:
    cols = list(df.columns)
    k = cols.index("результат, пунктов") + 1
    return [*cols[:k], "исход", "серия убыточных подряд", *cols[k:]]


def write_sheet(wb, m) -> tuple[str, int, dict]:
    df = m["df"]
    ws = wb.create_sheet(m["id"][:31])
    ws["A1"] = f"{m['strategy']} · {m['asset']} · {m['params']}"
    ws["A1"].font = font(size=12, bold=True)
    ws["A2"] = f"Параметры: {m['selection']}. {m['costs'].capitalize()}."
    ws["A3"] = ("Каждая строка — исход по порядку закрытия. Пункт — единица котировки "
                "(BR: 1 USD за баррель, NG: 1 USD за MMBtu); плюс — прибыль, минус — убыток.")
    for r in (2, 3):
        ws[f"A{r}"].font = font(color="52514E")
    cols = sheet_columns(df)
    letter = {c: get_column_letter(j) for j, c in enumerate(cols, start=1)}
    for j, c in enumerate(cols, start=1):
        cell = ws.cell(row=FIRST - 1, column=j, value=c)
        cell.font = font(bold=True)
        cell.fill = HEAD_FILL
        cell.border = THIN
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(j)].width = WIDTH.get(c, 12)
    ws.row_dimensions[FIRST - 1].height = 30
    plain = font()
    p = letter["результат, пунктов"]
    for i, row in enumerate(df.itertuples(index=False), start=FIRST):
        rec = dict(zip(df.columns, row, strict=True))
        for c in df.columns:
            v = rec[c]
            if pd.isna(v) or v == "":
                continue
            cell = ws[f"{letter[c]}{i}"]
            cell.value = v
            cell.font = plain
            cell.number_format = NUM_FMT.get(c, "@" if isinstance(v, str) else "General")
        ws[f"{letter['исход']}{i}"] = f'=IF({p}{i}>0,"прибыль",IF({p}{i}<0,"убыток","ноль"))'
        s = letter["серия убыточных подряд"]
        e = letter["вход"]
        # Переоценка открытой позиции на конец периода — не сделка: серию не продолжает и не обрывает.
        ws[f"{s}{i}"] = (f'=IF(AND({e}{i}<>"",{p}{i}<0),1,0)' if i == FIRST else
                         f'=IF({e}{i}="",{s}{i - 1},IF({p}{i}<0,{s}{i - 1}+1,0))')
        for c in ("исход", "серия убыточных подряд"):
            ws[f"{letter[c]}{i}"].font = plain
    ws.freeze_panes = f"A{FIRST}"
    ws.auto_filter.ref = f"A{FIRST - 1}:{get_column_letter(len(cols))}{FIRST + len(df) - 1}"
    return ws.title, FIRST + len(df) - 1, letter


def write_summary(ws, rows, title="Исходы сделок в пунктах: выгоднейшие параметры каждой стратегии"):
    ws["A1"] = title
    ws["A1"].font = font(size=14, bold=True)
    notes = ["Лист на каждую пару «стратегия × актив»: исходы по порядку, плюс — прибыль, минус — убыток.",
             "Пункт — единица котировки: BR — 1 USD за баррель, NG — 1 USD за MMBtu; шаг цены — 0.01 (BR) и "
             "0.001 (NG).",
             "Семь семейств: размер позиции меняется, поэтому пункты даны на единицу позиции, а итог — в USD "
             "на счёт 1 000 000. Скользящие средние: 1 контракт, пункты складываются.",
             "Все числа на этом листе — формулы по листам таблиц."]
    for k, t in enumerate(notes, start=2):
        ws[f"A{k}"] = t
        ws[f"A{k}"].font = font(color="52514E")
    head = ["Лист", "Стратегия", "Актив", "Параметры", "Исходов", "Прибыльных", "Убыточных", "Доля прибыльных",
            "Средний прибыльный, пунктов", "Средний убыточный, пунктов", "Лучший, пунктов", "Худший, пунктов",
            "Макс. убыточных подряд", "Итог IS, USD", "Итог OOS, USD", "Итог, USD"]
    widths = [26, 34, 7, 46, 9, 11, 11, 11, 14, 14, 12, 12, 12, 15, 15, 15]
    h = 7
    for j, (name, w) in enumerate(zip(head, widths, strict=True), start=1):
        c = ws.cell(row=h, column=j, value=name)
        c.font = font(bold=True)
        c.fill = HEAD_FILL
        c.border = THIN
        c.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.row_dimensions[h].height = 44
    for i, (m, sheet, last, L) in enumerate(rows, start=h + 1):
        rng = lambda col, sheet=sheet, last=last: f"'{sheet}'!${col}${FIRST}:${col}${last}"  # noqa: E731
        p, d, b, u = L["результат, пунктов"], L["вход"], L["период"], L["результат, USD"]
        ws[f"A{i}"] = sheet
        ws[f"A{i}"].hyperlink = f"#'{sheet}'!A1"
        ws[f"A{i}"].font = font(color="1C5CAB", underline="single")
        ws[f"B{i}"], ws[f"C{i}"], ws[f"D{i}"] = m["strategy"], m["asset"], m["params"]
        ws[f"E{i}"] = f'=COUNTIFS({rng(d)},"<>")'
        ws[f"F{i}"] = f'=COUNTIFS({rng(p)},">0",{rng(d)},"<>")'
        ws[f"G{i}"] = f'=COUNTIFS({rng(p)},"<0",{rng(d)},"<>")'
        ws[f"H{i}"] = f"=IF(E{i}>0,F{i}/E{i},0)"
        ws[f"I{i}"] = f'=IFERROR(AVERAGEIFS({rng(p)},{rng(p)},">0",{rng(d)},"<>"),0)'
        ws[f"J{i}"] = f'=IFERROR(AVERAGEIFS({rng(p)},{rng(p)},"<0",{rng(d)},"<>"),0)'
        ws[f"K{i}"] = f'=_xlfn.MAXIFS({rng(p)},{rng(d)},"<>")'
        ws[f"L{i}"] = f'=_xlfn.MINIFS({rng(p)},{rng(d)},"<>")'
        ws[f"M{i}"] = f"=MAX({rng(L['серия убыточных подряд'])})"
        ws[f"N{i}"] = f'=SUMIFS({rng(u)},{rng(b)},"IS")'
        ws[f"O{i}"] = f'=SUMIFS({rng(u)},{rng(b)},"OOS")'
        ws[f"P{i}"] = f"=SUM({rng(u)})"
        for col in "BCDEFGHIJKLMNOP":
            ws[f"{col}{i}"].font = font()
            ws[f"{col}{i}"].number_format = {"E": "0", "F": "0", "G": "0", "H": "0.0%", "M": "0"}.get(
                col, PT_FMT if col in "IJKL" else (USD if col in "NOP" else "@"))
    ws.freeze_panes = f"B{h + 1}"


def write_readme(items):
    lines = [
        "# Исходы сделок в пунктах: выгоднейшие параметры каждой стратегии", "",
        "Таблица на каждую пару «стратегия × актив»: все сделки по порядку закрытия и результат каждого "
        "исхода в пунктах — **плюс при прибыли, минус при убытке**. Те же таблицы одним файлом со статистикой "
        "на формулах: [`points.xlsx`](points.xlsx). **Только числа** — один столбец исходов в шагах цены, "
        "по файлу на стратегию и актив: [`sequence/`](sequence/README.md).", "",
        "* **Пункт** — единица биржевой котировки: у BR 1 USD за баррель, у NG 1 USD за MMBtu. Рядом — тот же "
        "результат в **шагах цены** (0.01 у BR, 0.001 у NG): если вы считаете пунктами минимальные шаги, "
        "смотрите этот столбец.",
        "* **Параметры.** Семь семейств — комбинация, выбранная оптимизатором (лучший Шарп в выборке, "
        "`docs/RESULTS_SEVEN_FAMILIES.md`). Скользящие средние — правило R1, максимальная прибыль в "
        "выборке, финальный этап 3.",
        "* **Издержки.** У семи семейств вычтены (~7 б.п. за оборот): пункты — по ценам исполнения минус "
        "комиссия. У скользящих средних сетка без издержек (условие задачи); столбец «при издержках 1 шаг» "
        "показывает исход за вычетом одного шага цены.",
        "* **Семь семейств.** Размер позиции меняется (целевая волатильность), уменьшение позиции — "
        "отдельный исход, как и в отчёте. Пункты даны на единицу позиции; `объём × пункты = USD`. Поэтому "
        "складывать пункты по строкам нельзя — нарастающий итог дан в USD на счёт 1 000 000. Портфельные "
        "стратегии (4 и 6) разложены по ногам NG и BR.",
        "* **Через ролл** — позиция пережила смену контракта: результат включает перекладку и не равен "
        "разности цен входа и выхода. **IS** — выборка, по которой выбирались параметры; **OOS** — проверка "
        "вне выборки.", "",
        "Сверки при выгрузке (`export_points.py` падает при расхождении): итоги вне выборки семи семейств "
        "совпадают с отчётом; цена выхода совпадает с биржевой котировкой дня с точностью до издержек; без "
        "ролла коэффициент склейки на входе и выходе один и тот же; `объём × пункты = USD`. У скользящих "
        "средних сумма совпадает со статистикой VectorBT финалиста.", "",
        "| Стратегия | Актив | Параметры | Исходов | Прибыльных | Средний +, пунктов | Средний −, пунктов "
        "| Лучший, пунктов | Худший, пунктов | Убыточных подряд, макс. | Итог IS, USD | Итог OOS, USD | Таблица |",
        "|---|---|---|---|---|---|---|---|---|---|---|---|---|",
    ]
    num = lambda v: f"{v:+,.0f}".replace(",", " ")  # noqa: E731
    for m in items:
        s = stats(m["df"])
        lines.append(f"| {m['strategy']} | {m['asset']} | {m['params']} | {s['n']} | {s['wins']} ({s['rate']:.0%}) "
                     f"| {s['avg_win']:+.4f} | {s['avg_loss']:+.4f} | {s['best']:+.4f} | {s['worst']:+.4f} "
                     f"| {s['streak']} | {num(s['usd_is'])} | {num(s['usd_oos'])} | [{m['id']}.csv]({m['id']}.csv) |")
    lines += ["", "Пересборка: `python export_points.py seven` (окружение пакета `exchange`), "
              "`python export_points.py ma`, `python export_points.py tables`."]
    (PTS / "README.md").write_text("\n".join(lines) + "\n")


def write_readme_full(items, pts: Path):
    lines = [
        "# Исходы сделок в пунктах: семь семейств, сквозной прогон", "",
        "Те же системы и параметры, что в [`../`](../README.md), но **одним прогоном по всей истории**: "
        "все сделки подряд, без разрыва между выборкой и проверкой. В прогонах отчёта OOS начинается без "
        "позиции и с пустыми индикаторами, а между периодами — эмбарго 10 дней; поэтому там нет сделок за "
        "несколько месяцев, а позиция на конец IS закрыта переоценкой. Здесь этих разрывов нет. Период "
        "сделки — по дате выхода: **IS**, **эмбарго** (10 дней между периодами) или **OOS**. Итоги OOS "
        "здесь не равны итогам отчёта: отчёт честно оценивает систему, запущенную с нуля в начале OOS.", "",
        "Тот же файл одним Excel со статистикой на формулах: [`points_full.xlsx`](points_full.xlsx). Только "
        "числа, по файлу на актив: [`../sequence/full/`](../sequence/full/README.md).", "",
        "| Стратегия | Актив | Параметры | Исходов | Прибыльных | Средний +, пунктов | Средний −, пунктов "
        "| Убыточных подряд, макс. | Итог IS, USD | Итог OOS, USD | Итог, USD | Таблица |",
        "|---|---|---|---|---|---|---|---|---|---|---|---|",
    ]
    num = lambda v: f"{v:+,.0f}".replace(",", " ")  # noqa: E731
    for m in items:
        s = stats(m["df"])
        lines.append(f"| {m['strategy']} | {m['asset']} | {m['params']} | {s['n']} | {s['wins']} ({s['rate']:.0%}) "
                     f"| {s['avg_win']:+.4f} | {s['avg_loss']:+.4f} | {s['streak']} | {num(s['usd_is'])} "
                     f"| {num(s['usd_oos'])} | {num(s['usd'])} | [{m['id']}.csv]({m['id']}.csv) |")
    lines += ["", "Пересборка: `python export_points.py seven-full` (окружение пакета `exchange`), "
              "`python export_points.py tables`, `python export_points.py sequence`."]
    (pts / "README.md").write_text("\n".join(lines) + "\n")


def main(pts: Path = PTS):
    items = tables(pts)
    wb = Workbook()
    summary = wb.active
    summary.title = "Сводка"
    rows = []
    for m in items:
        sheet, last, letter = write_sheet(wb, m)
        rows.append((m, sheet, last, letter))
    if pts == PTS:
        write_summary(summary, rows)
        out = PTS / "points.xlsx"
    else:
        write_summary(summary, rows, "Исходы сделок в пунктах: семь семейств, сквозной прогон по всей истории")
        out = pts / "points_full.xlsx"
    wb.save(out)
    if pts == PTS:
        write_readme(items)
    else:
        write_readme_full(items, pts)
    print(f"{len(items)} таблиц -> {out} ({out.stat().st_size / 1e6:.1f} МБ) и README.md")


if __name__ == "__main__":
    main()
